"""Administración: sucursales, usuarios, clientes, notificaciones internas, auditoría y reportes."""
import csv
import io
from datetime import datetime

from flask import Blueprint, Response, g, jsonify, request, send_file
from sqlalchemy import func, or_, select, update

from ..errors import AppError, NoEncontrado
from ..extensions import db
from ..models import (Auditoria, Cliente, Devolucion, Inventario, MovimientoInventario, Notificacion, Pago, Pedido,
                      Rol, Sucursal, SucursalZona, Usuario, Variante)
from ..schemas import SucursalSchema, UsuarioSchema, cargar
from ..security.auth import (filtro_sucursal, hash_password, requiere_usuario, sucursales_permitidas, validar_fortaleza,
                             verificar_alcance)
from ..services import reporte_service as rep
from ..services.comun import auditar, paginar
from . import serial
from .util import arg_int, cuerpo, obtener, pagina

bp = Blueprint("admin_sistema", __name__, url_prefix="/api/admin")


# ---------------------------------------------------------------- sucursales

@bp.get("/sucursales")
@requiere_usuario("sucursal.ver")
def sucursales():
    q = select(Sucursal).order_by(Sucursal.id)
    return jsonify([serial.sucursal(s, completo=True) for s in db.session.scalars(q)])


def _aplicar_sucursal(s: Sucursal, d: dict):
    zonas = d.pop("zonas", [])
    if d.get("responsable_id") and not db.session.get(Usuario, d["responsable_id"]):
        raise NoEncontrado("Responsable no encontrado.")
    for k, v in d.items():
        setattr(s, k, v)
    s.zonas = [SucursalZona(cp_prefijo=str(z["cp_prefijo"]), prioridad=int(z.get("prioridad", 1))) for z in zonas]


@bp.post("/sucursales")
@requiere_usuario("sucursal.gestionar")
def crear_sucursal():
    d = cargar(SucursalSchema, cuerpo())
    s = Sucursal()
    _aplicar_sucursal(s, d)
    db.session.add(s)
    db.session.flush()
    auditar("sucursal.crear", "sucursal", s.id, despues=d, sucursal_id=s.id)
    db.session.commit()
    return jsonify(serial.sucursal(s, True)), 201


@bp.put("/sucursales/<int:sucursal_id>")
@requiere_usuario("sucursal.gestionar")
def editar_sucursal(sucursal_id):
    s = obtener(Sucursal, sucursal_id, "Sucursal no encontrada.")
    antes = serial.sucursal(s, True)
    d = cargar(SucursalSchema, cuerpo())
    _aplicar_sucursal(s, d)
    auditar("sucursal.editar", "sucursal", s.id, antes=antes, despues=d, sucursal_id=s.id)
    db.session.commit()
    return jsonify(serial.sucursal(s, True))


# ---------------------------------------------------------------- usuarios y roles

@bp.get("/roles")
@requiere_usuario("usuario.gestionar")
def roles():
    return jsonify([{"codigo": r.codigo, "nombre": r.nombre, "descripcion": r.descripcion, "alcance_global": r.alcance_global,
                     "permisos": sorted(p.codigo for p in r.permisos)} for r in db.session.scalars(select(Rol).order_by(Rol.id))])


@bp.get("/usuarios")
@requiere_usuario("usuario.gestionar")
def usuarios():
    q = select(Usuario).order_by(Usuario.nombre)
    if texto := request.args.get("q"):
        q = q.where(or_(Usuario.nombre.ilike(f"%{texto}%"), Usuario.correo.ilike(f"%{texto}%")))
    return jsonify([serial.usuario(u) for u in db.session.scalars(q).unique()])


def _aplicar_usuario(u: Usuario, d: dict):
    rol = db.session.scalars(select(Rol).where(Rol.codigo == d["rol"])).first()
    if not rol:
        raise AppError("Rol inválido.")
    sucs = db.session.scalars(select(Sucursal).where(Sucursal.id.in_(d["sucursales"] or [-1]))).all()
    if len(sucs) != len(set(d["sucursales"])):
        raise NoEncontrado("Alguna sucursal no existe.")
    if not rol.alcance_global and not sucs:
        raise AppError("Los roles de sucursal requieren al menos una sucursal asignada.")
    if u.id == g.usuario.id and (not d["activo"] or rol.codigo != "admin") and g.usuario.rol.codigo == "admin":
        raise AppError("No puedes desactivarte ni quitarte el rol de administrador a ti mismo.")
    u.correo, u.nombre, u.rol, u.activo, u.sucursales = d["correo"].lower(), d["nombre"], rol, d["activo"], sucs
    if d.get("password"):
        validar_fortaleza(d["password"])
        u.hash_password = hash_password(d["password"])


@bp.post("/usuarios")
@requiere_usuario("usuario.gestionar")
def crear_usuario():
    d = cargar(UsuarioSchema, cuerpo())
    if not d.get("password"):
        raise AppError("Asigna una contraseña inicial.")
    u = Usuario()
    _aplicar_usuario(u, d)
    db.session.add(u)
    db.session.flush()
    auditar("usuario.crear", "usuario", u.id, despues={"correo": u.correo, "rol": d["rol"], "sucursales": d["sucursales"]})
    db.session.commit()
    return jsonify(serial.usuario(u)), 201


@bp.put("/usuarios/<int:usuario_id>")
@requiere_usuario("usuario.gestionar")
def editar_usuario(usuario_id):
    u = obtener(Usuario, usuario_id, "Usuario no encontrado.")
    antes = serial.usuario(u)
    d = cargar(UsuarioSchema, cuerpo())
    _aplicar_usuario(u, d)
    auditar("usuario.editar", "usuario", u.id, antes={k: antes[k] for k in ("rol", "sucursales", "activo")},
            despues={"rol": d["rol"], "sucursales": d["sucursales"], "activo": d["activo"],
                     "password_cambiada": bool(d.get("password"))})
    db.session.commit()
    return jsonify(serial.usuario(u))


@bp.post("/usuarios/<int:usuario_id>/desbloquear")
@requiere_usuario("usuario.gestionar")
def desbloquear(usuario_id):
    u = obtener(Usuario, usuario_id, "Usuario no encontrado.")
    u.bloqueado_hasta, u.intentos_fallidos = None, 0
    auditar("usuario.desbloquear", "usuario", u.id)
    db.session.commit()
    return jsonify(serial.usuario(u))


# ---------------------------------------------------------------- clientes

@bp.get("/clientes")
@requiere_usuario("cliente.ver")
def clientes():
    q = select(Cliente).order_by(Cliente.id.desc())
    if texto := request.args.get("q"):
        q = q.where(or_(Cliente.nombre.ilike(f"%{texto}%"), Cliente.correo.ilike(f"%{texto}%"),
                        Cliente.razon_social.ilike(f"%{texto}%")))
    if tipo := request.args.get("tipo"):
        q = q.where(Cliente.tipo == tipo)
    return jsonify(paginar(q, *pagina(), serial.cliente))


@bp.get("/clientes/<int:cliente_id>")
@requiere_usuario("cliente.ver")
def cliente(cliente_id):
    c = obtener(Cliente, cliente_id, "Cliente no encontrado.")
    pedidos = db.session.scalars(select(Pedido).where(Pedido.cliente_id == c.id).order_by(Pedido.id.desc()).limit(50)).unique()
    pagos = db.session.scalars(select(Pago).join(Pedido).where(Pedido.cliente_id == c.id).order_by(Pago.id.desc()).limit(50)).unique()
    d = serial.cliente(c)
    d["direcciones"] = [serial.direccion(x) for x in c.direcciones if x.activo]
    d["pedidos"] = [serial.pedido_resumen(p) for p in pedidos]
    d["pagos"] = [serial.pago(p) for p in pagos]
    return jsonify(d)


# ---------------------------------------------------------------- notificaciones del personal

@bp.get("/notificaciones")
@requiere_usuario()
def notificaciones():
    q = select(Notificacion).where(Notificacion.usuario_id == g.usuario.id).order_by(Notificacion.id.desc())
    if request.args.get("no_leidas") in {"1", "true"}:
        q = q.where(Notificacion.leida.is_(False))
    datos = paginar(q, *pagina(), serial.notificacion)
    datos["no_leidas"] = db.session.scalar(select(func.count()).select_from(Notificacion).where(
        Notificacion.usuario_id == g.usuario.id, Notificacion.leida.is_(False)))
    return jsonify(datos)


@bp.post("/notificaciones/leer")
@requiere_usuario()
def leer():
    ids = cuerpo().get("ids")
    q = update(Notificacion).where(Notificacion.usuario_id == g.usuario.id)
    if ids:
        q = q.where(Notificacion.id.in_([int(i) for i in ids]))
    db.session.execute(q.values(leida=True))
    db.session.commit()
    return jsonify({"ok": True})


# ---------------------------------------------------------------- auditoría

@bp.get("/auditoria")
@requiere_usuario("auditoria.ver")
def auditoria():
    q = select(Auditoria).order_by(Auditoria.id.desc())
    if entidad := request.args.get("entidad"):
        q = q.where(Auditoria.entidad == entidad)
    if (eid := arg_int("entidad_id")) is not None:
        q = q.where(Auditoria.entidad_id == eid)
    if accion := request.args.get("accion"):
        q = q.where(Auditoria.accion.ilike(f"%{accion}%"))
    if (sid := arg_int("sucursal_id")) is not None:
        q = q.where(Auditoria.sucursal_id == sid)
    cond = filtro_sucursal(Auditoria.sucursal_id)
    if cond is not None:
        q = q.where(cond)
    return jsonify(paginar(q, *pagina(), serial.auditoria))


# ---------------------------------------------------------------- reportes

@bp.get("/reportes/dashboard")
@requiere_usuario("reporte.ver")
def dashboard():
    sid = arg_int("sucursal_id")
    if sid:
        verificar_alcance(sid)
    return jsonify(rep.dashboard(request.args.get("desde"), request.args.get("hasta"), sid))


@bp.get("/reportes/sucursales")
@requiere_usuario("reporte.ver")
def comparativa():
    return jsonify(rep.comparativa_sucursales(request.args.get("desde"), request.args.get("hasta")))


@bp.get("/reportes/existencias-bajas")
@requiere_usuario("inventario.ver")
def bajas():
    return jsonify(rep.existencias_bajas(arg_int("sucursal_id"), limite=500))


def _exportar(nombre: str, columnas: list[str], filas: list[list]):
    formato = request.args.get("formato", "csv")
    sello = datetime.now().strftime("%Y%m%d-%H%M")
    if formato == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        wb = Workbook()
        ws = wb.active
        ws.title = nombre[:30]
        ws.append(columnas)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="142B4A")
        for fila in filas:
            ws.append(fila)
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(45, max(12, *(len(str(c.value or "")) + 2 for c in col)))
        ws.freeze_panes = "A2"
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return send_file(buf, as_attachment=True, download_name=f"{nombre}-{sello}.xlsx",
                         mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(columnas)
    for fila in filas:
        w.writerow([("'" + v) if isinstance(v, str) and v[:1] in "=+-@" else v for v in fila])  # evita inyección de fórmulas
    return Response("﻿" + buf.getvalue(), mimetype="text/csv; charset=utf-8",
                    headers={"Content-Disposition": f"attachment; filename={nombre}-{sello}.csv"})


@bp.get("/reportes/exportar/<reporte>")
@requiere_usuario("reporte.ver")
def exportar(reporte):
    ini, fin = rep.rango(request.args.get("desde"), request.args.get("hasta"))
    permitidas = sucursales_permitidas()
    if reporte == "pedidos":
        q = select(Pedido).where(Pedido.creado.between(ini, fin)).order_by(Pedido.id)
        alcance = rep._pedidos_en_alcance(arg_int("sucursal_id"))
        if alcance is not None:
            q = q.where(Pedido.id.in_(alcance))
        filas = [[p.folio, p.creado.strftime("%Y-%m-%d %H:%M"), p.canal, p.cliente.razon_social or p.cliente.nombre,
                  p.estado, p.modalidad, p.metodo_pago, float(p.subtotal), float(p.descuento), float(p.costo_envio),
                  float(p.total), ", ".join(sorted({a.sucursal.nombre for a in p.asignaciones}))]
                 for p in db.session.scalars(q).unique()]
        return _exportar("pedidos", ["Folio", "Fecha", "Canal", "Cliente", "Estado", "Modalidad", "Método de pago",
                                     "Subtotal", "Descuento", "Envío", "Total", "Sucursales"], filas)
    if reporte == "inventario":
        q = select(Inventario).join(Variante).order_by(Inventario.sucursal_id, Variante.sku)
        if permitidas is not None:
            q = q.where(Inventario.sucursal_id.in_(permitidas or {-1}))
        if sid := arg_int("sucursal_id"):
            q = q.where(Inventario.sucursal_id == sid)
        filas = [[i.sucursal.nombre, i.variante.sku, i.variante.producto.nombre, i.variante.talla, i.variante.color,
                  i.fisico, i.reservado, i.comprometido, i.en_transito, i.disponible, i.minimo]
                 for i in db.session.scalars(q).unique()]
        return _exportar("inventario", ["Sucursal", "SKU", "Producto", "Talla", "Color", "Físico", "Reservado",
                                        "Comprometido", "En tránsito", "Disponible", "Mínimo"], filas)
    if reporte == "movimientos":
        q = select(MovimientoInventario).where(MovimientoInventario.creado.between(ini, fin)).order_by(MovimientoInventario.id)
        if permitidas is not None:
            q = q.where(MovimientoInventario.sucursal_id.in_(permitidas or {-1}))
        if sid := arg_int("sucursal_id"):
            q = q.where(MovimientoInventario.sucursal_id == sid)
        nombres = {s.id: s.nombre for s in db.session.scalars(select(Sucursal))}
        filas = [[m.creado.strftime("%Y-%m-%d %H:%M"), nombres.get(m.sucursal_id), m.variante.sku, m.tipo, m.cantidad,
                  m.saldo_resultante, m.motivo, m.usuario.nombre if m.usuario else "Sistema"]
                 for m in db.session.scalars(q).unique()]
        return _exportar("movimientos", ["Fecha", "Sucursal", "SKU", "Tipo", "Cantidad", "Saldo", "Motivo", "Usuario"], filas)
    if reporte == "pagos":
        q = select(Pago).where(Pago.creado.between(ini, fin)).order_by(Pago.id)
        alcance = rep._pedidos_en_alcance(arg_int("sucursal_id"))
        if alcance is not None:
            q = q.where(Pago.pedido_id.in_(alcance))
        filas = [[p.folio, p.creado.strftime("%Y-%m-%d %H:%M"), p.pedido.folio, p.metodo, p.estado, float(p.monto),
                  float(p.monto_reembolsado), p.referencia or "", p.validado_en.strftime("%Y-%m-%d %H:%M") if p.validado_en else ""]
                 for p in db.session.scalars(q).unique()]
        return _exportar("pagos", ["Folio", "Fecha", "Pedido", "Método", "Estado", "Monto", "Reembolsado",
                                   "Referencia", "Validado"], filas)
    if reporte == "devoluciones":
        q = select(Devolucion).where(Devolucion.creado.between(ini, fin)).order_by(Devolucion.id)
        alcance = rep._pedidos_en_alcance(arg_int("sucursal_id"))
        if alcance is not None:
            q = q.where(Devolucion.pedido_id.in_(alcance))
        filas = [[d.folio, d.creado.strftime("%Y-%m-%d"), d.pedido.folio, d.motivo, d.estado,
                  sum(x.cantidad for x in d.detalles), float(d.monto_reembolso)] for d in db.session.scalars(q).unique()]
        return _exportar("devoluciones", ["Folio", "Fecha", "Pedido", "Motivo", "Estado", "Unidades", "Monto"], filas)
    if reporte == "sucursales":
        datos = rep.comparativa_sucursales(request.args.get("desde"), request.args.get("hasta"))
        filas = [[d["sucursal"], d["pedidos"], d["unidades_asignadas"], d["unidades_despachadas"], d["faltantes"],
                  float(d["importe_surtido"]), d["existencia_unidades"], float(d["existencia_valor"]), d["existencias_bajas"]]
                 for d in datos]
        return _exportar("comparativa-sucursales", ["Sucursal", "Pedidos", "Unidades asignadas", "Despachadas",
                                                    "Faltantes", "Importe surtido", "Existencia (u)", "Valor existencia",
                                                    "Existencias bajas"], filas)
    raise NoEncontrado("Reporte no disponible.")


@bp.get("/pedidos/<int:pedido_id>/pdf")
@requiere_usuario("pedido.ver")
def pdf_pedido(pedido_id):
    from ..services.documentos import pdf_pedido as generar
    from ..services.pedido_service import sucursales_de
    p = obtener(Pedido, pedido_id, "Pedido no encontrado.")
    verificar_alcance(sucursales_de(p) or None)
    return send_file(io.BytesIO(generar(p)), as_attachment=True, download_name=f"{p.folio}.pdf", mimetype="application/pdf")
